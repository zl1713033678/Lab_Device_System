import asyncio
import csv
import io
import json
import os
import shutil
import struct
import zlib
from abc import ABC, abstractmethod
from datetime import datetime, timedelta
from enum import Enum
from typing import Any, Dict, List, Optional

from fastapi import FastAPI, HTTPException, WebSocket, WebSocketDisconnect
from fastapi.responses import Response
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel
from sqlalchemy import Column, DateTime, Integer, String, create_engine, inspect, text
from sqlalchemy.orm import declarative_base, sessionmaker
import pandas as pd

# ------------------------------------------------------------------
# 1. 数据库配置与自动迁移
# ------------------------------------------------------------------
DATABASE_URL = "sqlite:///./lab_devices.db"
engine = create_engine(DATABASE_URL, connect_args={"check_same_thread": False})
SessionLocal = sessionmaker(autocommit=False, autoflush=False, bind=engine)
Base = declarative_base()


def check_and_migrate_db():
    """系统级数据库模式在线热迁移：安全兼容旧表并无缝扩容字段"""
    inspector = inspect(engine)
    if "devices" in inspector.get_table_names():
        columns = [c["name"] for c in inspector.get_columns("devices")]
        with engine.begin() as conn:
            if "room_name" not in columns:
                conn.execute(
                    text("ALTER TABLE devices ADD COLUMN room_name VARCHAR DEFAULT '测试实验室 1'")
                )
            if "pos_x" not in columns:
                conn.execute(
                    text("ALTER TABLE devices ADD COLUMN pos_x VARCHAR DEFAULT '50%'")
                )
                if "pos_left" in columns:
                    conn.execute(text("UPDATE devices SET pos_x = pos_left"))
            if "pos_y" not in columns:
                conn.execute(
                    text("ALTER TABLE devices ADD COLUMN pos_y VARCHAR DEFAULT '50%'")
                )
                if "pos_top" in columns:
                    conn.execute(text("UPDATE devices SET pos_y = pos_top"))

    if "rooms" in inspector.get_table_names():
        columns = [c["name"] for c in inspector.get_columns("rooms")]
        with engine.begin() as conn:
            if "pos_x" not in columns:
                conn.execute(
                    text("ALTER TABLE rooms ADD COLUMN pos_x VARCHAR DEFAULT '50%'")
                )
            if "pos_y" not in columns:
                conn.execute(
                    text("ALTER TABLE rooms ADD COLUMN pos_y VARCHAR DEFAULT '50%'")
                )

    if "device_usage_logs" in inspector.get_table_names():
        columns = [c["name"] for c in inspector.get_columns("device_usage_logs")]
        with engine.begin() as conn:
            if "raw_data" not in columns:
                conn.execute(
                    text("ALTER TABLE device_usage_logs ADD COLUMN raw_data TEXT")
                )
                if "raw_payload" in columns:
                    conn.execute(text("UPDATE device_usage_logs SET raw_data = raw_payload"))


# ------------------------------------------------------------------
# 2. 数据库表定义
# ------------------------------------------------------------------
class RoomModel(Base):
    """实验室房间/区域表"""
    __tablename__ = "rooms"
    id = Column(Integer, primary_key=True, index=True)
    name = Column(String, unique=True, index=True)
    pos_x = Column(String, default="50%")
    pos_y = Column(String, default="50%")


class DeviceStatus(str, Enum):
    IDLE = "IDLE"
    IN_USE = "IN_USE"
    FAULT = "FAULT"


class DeviceModel(Base):
    __tablename__ = "devices"
    id = Column(Integer, primary_key=True, index=True)
    device_code = Column(String, unique=True, index=True)
    name = Column(String)
    room_name = Column(String, default="302")
    status = Column(String, default=DeviceStatus.IDLE.value)
    pos_x = Column(String, default="50%")
    pos_y = Column(String, default="50%")


class DeviceUsageLogModel(Base):
    """不可篡改的系统审计日志表 (Append-Only)"""
    __tablename__ = "device_usage_logs"
    id = Column(Integer, primary_key=True, index=True)
    device_code = Column(String, index=True)
    previous_status = Column(String)
    new_status = Column(String)
    operator = Column(String, default="SYSTEM_MANUAL")
    source_type = Column(String, default="MANUAL_WEB")
    raw_data = Column(String, nullable=True)
    timestamp = Column(DateTime, default=datetime.now)


class SystemSettingModel(Base):
    """系统全局配置表"""
    __tablename__ = "system_settings"
    key = Column(String, primary_key=True)
    value = Column(String)


# 执行数据库迁移并初始化表格
check_and_migrate_db()
Base.metadata.create_all(bind=engine)


# ------------------------------------------------------------------
# 3. WebSocket 实时连接管理器 (ConnectionManager)
# ------------------------------------------------------------------
class ConnectionManager:
    """管理在线 WebSocket 长连接，实现状态变更毫秒级广播"""
    def __init__(self):
        self.active_connections: List[WebSocket] = []

    async def connect(self, websocket: WebSocket):
        await websocket.accept()
        self.active_connections.append(websocket)

    def disconnect(self, websocket: WebSocket):
        if websocket in self.active_connections:
            self.active_connections.remove(websocket)

    async def broadcast(self, message: dict):
        """事件驱动推送"""
        for connection in list(self.active_connections):
            try:
                await connection.send_json(message)
            except Exception:
                self.disconnect(connection)


manager = ConnectionManager()


# ------------------------------------------------------------------
# 4. Pydantic 遥测与管理报文模型定义
# ------------------------------------------------------------------
class TelemetryChannel(BaseModel):
    device_code: str
    status: Optional[DeviceStatus] = None
    state: Optional[DeviceStatus] = None
    current: Optional[float] = None
    voltage: Optional[float] = None
    operator: Optional[str] = "SYSTEM"
    source_type: Optional[str] = "ESP32"
    raw_data: Optional[Any] = None


class TelemetryIngestPayload(BaseModel):
    device_code: Optional[str] = None
    status: Optional[DeviceStatus] = None
    state: Optional[DeviceStatus] = None
    current: Optional[float] = None
    voltage: Optional[float] = None
    operator: Optional[str] = "SYSTEM"
    source_type: Optional[str] = "MANUAL_WEB"
    raw_data: Optional[Any] = None
    channels: Optional[List[TelemetryChannel]] = None


class StatusChangeRequest(BaseModel):
    new_status: DeviceStatus
    operator: Optional[str] = "MANUAL_WEB"


class CreateDeviceRequest(BaseModel):
    name: str
    device_code: str
    room_name: Optional[str] = "302"
    pos_x: Optional[str] = None
    pos_y: Optional[str] = None


class CreateRoomRequest(BaseModel):
    name: str


class UpdateRoomNameRequest(BaseModel):
    new_name: str


class UpdateRoomPositionRequest(BaseModel):
    pos_x: str
    pos_y: str


class UpdateDeviceRequest(BaseModel):
    name: str
    room_name: str
    pos_x: Optional[str] = None
    pos_y: Optional[str] = None


# ------------------------------------------------------------------
# 5. 数据源适配器与防抖降噪过滤器
# ------------------------------------------------------------------
class BaseTelemetryAdapter(ABC):
    @abstractmethod
    def adapt(self, payload: TelemetryChannel, previous_status: str) -> DeviceStatus:
        pass


class DefaultTelemetryAdapter(BaseTelemetryAdapter):
    def adapt(self, payload: TelemetryChannel, previous_status: str) -> DeviceStatus:
        # 如果前一状态为 FAULT (故障)，且本次数据来源并非 Web 端人工点击操作 (MANUAL_WEB)，
        # 则强制锁定 FAULT 故障锁定状态，即使有电流数据也不自动切回其它状态，直到人工主动点击切换
        if previous_status == DeviceStatus.FAULT.value and (payload.source_type or "").upper() != "MANUAL_WEB":
            return DeviceStatus.FAULT

        if payload.current is not None:
            curr = payload.current
            if curr > 0.5:
                return DeviceStatus.IN_USE
            elif curr <= 0.05:
                return DeviceStatus.IDLE
            else:
                if previous_status == DeviceStatus.IN_USE.value:
                    return DeviceStatus.IN_USE
                return DeviceStatus.IDLE

        target = payload.status or payload.state
        if target is not None:
            return target
        return DeviceStatus(previous_status)


class ManualWebAdapter(BaseTelemetryAdapter):
    def adapt(self, payload: TelemetryChannel, previous_status: str) -> DeviceStatus:
        target = payload.status or payload.state
        if target is not None:
            return target
        return DeviceStatus(previous_status)


class ESP32MQTTAdapter(BaseTelemetryAdapter):
    def adapt(self, payload: TelemetryChannel, previous_status: str) -> DeviceStatus:
        return DefaultTelemetryAdapter().adapt(payload, previous_status)


class PLCModbusAdapter(BaseTelemetryAdapter):
    def adapt(self, payload: TelemetryChannel, previous_status: str) -> DeviceStatus:
        return DefaultTelemetryAdapter().adapt(payload, previous_status)


class TelemetryAdapterFactory:
    _adapters = {
        "MANUAL_WEB": ManualWebAdapter(),
        "ESP32_MQTT": ESP32MQTTAdapter(),
        "ESP32": ESP32MQTTAdapter(),
        "PLC_MODBUS": PLCModbusAdapter(),
    }

    @classmethod
    def get_adapter(cls, source_type: str) -> BaseTelemetryAdapter:
        st = source_type.upper() if source_type else "MANUAL_WEB"
        return cls._adapters.get(st, DefaultTelemetryAdapter())


# ------------------------------------------------------------------
# 6. 硬件离线/断电心跳监控器
# ------------------------------------------------------------------
device_last_seen: Dict[str, datetime] = {}
device_offline_notified: set = set()


async def heartbeat_checker_loop():
    """后台硬件断电检测异步任务：10秒未收到遥测，判定为采样设备断电，自动切至待机并广播提醒"""
    while True:
        await asyncio.sleep(3)
        try:
            now = datetime.now()
            db = SessionLocal()
            try:
                devices = db.query(DeviceModel).all()
                for dev in devices:
                    code = dev.device_code
                    last_seen = device_last_seen.get(code)
                    if last_seen and (now - last_seen).total_seconds() > 10:
                        if code not in device_offline_notified:
                            device_offline_notified.add(code)
                            old_status = dev.status
                            # 如果之前的状态不是 FAULT，自动将其恢复为待机就绪 (IDLE)
                            if old_status != DeviceStatus.FAULT.value and old_status != DeviceStatus.IDLE.value:
                                dev.status = DeviceStatus.IDLE.value
                                log = DeviceUsageLogModel(
                                    device_code=code,
                                    previous_status=old_status,
                                    new_status=DeviceStatus.IDLE.value,
                                    operator="SYSTEM_TIMEOUT",
                                    source_type="HARDWARE_OFFLINE",
                                    raw_data=json.dumps({"reason": "采样设备断电/通信中断超时"}),
                                    timestamp=now,
                                )
                                db.add(log)
                                db.commit()

                            await manager.broadcast({
                                "type": "device_offline",
                                "device_code": code,
                                "name": dev.name,
                                "previous_status": old_status,
                                "new_status": dev.status,
                                "message": f"⚡ 采样硬件设备 [{code}] 突然断电离线！状态已自动切换为待机就绪",
                                "timestamp": now.strftime("%Y-%m-%d %H:%M:%S"),
                            })
            finally:
                db.close()
        except Exception as e:
            print(f"硬件心跳检测异常: {e}")


# ------------------------------------------------------------------
# 7. 核心数据写入与 WebSocket 毫秒广播管道
# ------------------------------------------------------------------
async def process_telemetry_ingest_single(db, ch: TelemetryChannel) -> dict:
    # 记录硬件最后活跃时间，并重置离线标记
    device_last_seen[ch.device_code] = datetime.now()
    if ch.device_code in device_offline_notified:
        device_offline_notified.remove(ch.device_code)

    device = db.query(DeviceModel).filter(DeviceModel.device_code == ch.device_code).first()
    
    if not device:
        unassigned_room = db.query(RoomModel).filter(RoomModel.name == "未分配实验室").first()
        if not unassigned_room:
            unassigned_room = RoomModel(name="未分配实验室", pos_x="50%", pos_y="50%")
            db.add(unassigned_room)
            db.commit()

        device = DeviceModel(
            device_code=ch.device_code.strip(),
            name=f"新硬件_{ch.device_code.strip()}",
            room_name="未分配实验室",
            status=DeviceStatus.IDLE.value,
            pos_x="50%",
            pos_y="50%",
        )
        db.add(device)
        db.commit()
        db.refresh(device)
        
        await manager.broadcast({
            "type": "device_discovered",
            "message": f"⚡ 检测到新 ESP32 硬件 [{device.device_code}] 接入网络，已自动录入未分配实验室！",
            "device": {
                "id": device.id,
                "device_code": device.device_code,
                "name": device.name,
                "room_name": device.room_name,
                "status": device.status,
                "pos_x": device.pos_x,
                "pos_y": device.pos_y,
            },
            "timestamp": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        })

    old_status = device.status
    source_type = ch.source_type or "PLC_MODBUS"
    adapter = TelemetryAdapterFactory.get_adapter(source_type)
    
    resolved_status = adapter.adapt(ch, old_status)
    new_status = resolved_status.value

    curr = ch.current
    volt = ch.voltage
    if curr is None and ch.raw_data and isinstance(ch.raw_data, dict):
        curr = ch.raw_data.get("current")
        if curr is None and "adc_diff" in ch.raw_data:
            curr = round(ch.raw_data["adc_diff"] * 0.008, 2)
    if volt is None and ch.raw_data and isinstance(ch.raw_data, dict):
        volt = ch.raw_data.get("voltage", 220.0)

    if old_status != new_status:
        device.status = new_status
        raw_json_str = json.dumps(ch.raw_data, ensure_ascii=False) if ch.raw_data else None

        log = DeviceUsageLogModel(
            device_code=ch.device_code,
            previous_status=old_status,
            new_status=new_status,
            operator=ch.operator or "SYSTEM",
            source_type=source_type,
            raw_data=raw_json_str,
            timestamp=datetime.now(),
        )
        db.add(log)
        db.commit()
        db.refresh(device)

        event_payload = {
            "type": "status_update",
            "device_code": device.device_code,
            "name": device.name,
            "previous_status": old_status,
            "new_status": new_status,
            "current": curr if curr is not None else 0.0,
            "voltage": volt if volt is not None else 220.0,
            "room_name": device.room_name or "302",
            "source_type": source_type,
            "operator": ch.operator or "SYSTEM",
            "timestamp": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        }
        await manager.broadcast(event_payload)
    else:
        event_payload = {
            "type": "telemetry_update",
            "device_code": device.device_code,
            "status": device.status,
            "current": curr if curr is not None else 0.0,
            "voltage": volt if volt is not None else 220.0,
            "raw_data": ch.raw_data,
            "timestamp": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        }
        await manager.broadcast(event_payload)

    return {
        "device_code": ch.device_code,
        "previous_status": old_status,
        "new_status": new_status,
        "room_name": device.room_name,
        "source_type": source_type,
    }


# ------------------------------------------------------------------
# 7. 履历核心计算算法
# ------------------------------------------------------------------
def format_duration(seconds: float) -> str:
    total_sec = int(seconds)
    hours = total_sec // 3600
    minutes = (total_sec % 3600) // 60
    secs = total_sec % 60
    if hours > 0:
        return f"{hours}小时 {minutes}分钟"
    elif minutes > 0:
        return f"{minutes}分钟 {secs}秒"
    else:
        return f"{secs}秒"


def calculate_device_usage_sessions(db, device_code: str, device_name: str, room_name: str) -> List[dict]:
    logs = (
        db.query(DeviceUsageLogModel)
        .filter(DeviceUsageLogModel.device_code == device_code)
        .order_by(DeviceUsageLogModel.timestamp.asc())
        .all()
    )

    sessions = []
    active_session = None

    for log in logs:
        prev_s = log.previous_status
        new_s = log.new_status
        
        if new_s == DeviceStatus.IN_USE.value and prev_s != DeviceStatus.IN_USE.value:
            if active_session:
                active_session["end_time"] = log.timestamp
                active_session["duration"] = max(0.0, (log.timestamp - active_session["start_time"]).total_seconds())
                active_session["formatted_duration"] = format_duration(active_session["duration"])
                active_session["status"] = "COMPLETED"
                sessions.append(active_session)
                
            active_session = {
                "device_code": device_code,
                "name": device_name,
                "room_name": room_name,
                "start_time": log.timestamp,
                "end_time": None,
                "duration": 0.0,
                "formatted_duration": "运行中",
                "operator": log.operator,
                "source_type": log.source_type,
                "start_raw": log.raw_data,
                "end_raw": None,
                "status": "RUNNING"
            }
        
        elif prev_s == DeviceStatus.IN_USE.value and new_s != DeviceStatus.IN_USE.value:
            if active_session:
                active_session["end_time"] = log.timestamp
                active_session["duration"] = max(0.0, (log.timestamp - active_session["start_time"]).total_seconds())
                active_session["formatted_duration"] = format_duration(active_session["duration"])
                active_session["end_raw"] = log.raw_data
                active_session["status"] = "COMPLETED"
                sessions.append(active_session)
                active_session = None

    device = db.query(DeviceModel).filter(DeviceModel.device_code == device_code).first()
    if device and device.status == DeviceStatus.IN_USE.value:
        if active_session:
            now = datetime.now()
            active_session["end_time"] = now
            active_session["duration"] = max(0.0, (now - active_session["start_time"]).total_seconds())
            active_session["formatted_duration"] = "运行中"
            sessions.append(active_session)
        else:
            start_t = datetime.now() - timedelta(hours=1)
            last_log = (
                db.query(DeviceUsageLogModel)
                .filter(DeviceUsageLogModel.device_code == device_code, DeviceUsageLogModel.new_status == DeviceStatus.IN_USE.value)
                .order_by(DeviceUsageLogModel.timestamp.desc())
                .first()
            )
            if last_log:
                start_t = last_log.timestamp
            sessions.append({
                "device_code": device_code,
                "name": device_name,
                "room_name": room_name,
                "start_time": start_t,
                "end_time": datetime.now(),
                "duration": max(0.0, (datetime.now() - start_t).total_seconds()),
                "formatted_duration": "运行中",
                "operator": "SYSTEM",
                "source_type": "UNKNOWN",
                "start_raw": None,
                "end_raw": None,
                "status": "RUNNING"
            })
    elif active_session:
        active_session["end_time"] = active_session["start_time"] + timedelta(minutes=1)
        active_session["duration"] = 60.0
        active_session["formatted_duration"] = format_duration(60.0)
        active_session["status"] = "COMPLETED"
        sessions.append(active_session)

    return sessions


def generate_default_cad_floorplan(filepath: str, width: int = 1200, height: int = 800):
    """用纯 Python (zlib+struct) 生成一张暗色科技风 CAD 底图"""
    rows = []
    for y in range(height):
        row = bytearray([0])
        for x in range(width):
            r, g, b = 20, 24, 30
            if x % 40 == 0 or y % 40 == 0:
                r, g, b = 32, 42, 54
            if x < 15 or x > width - 15 or y < 15 or y > height - 15:
                r, g, b = 0, 120, 180

            x1_min, x1_max = int(width * 0.10), int(width * 0.48)
            y1_min, y1_max = int(height * 0.12), int(height * 0.52)
            if (x1_min <= x <= x1_max and (y == y1_min or y == y1_max)) or \
               (y1_min <= y <= y1_max and (x == x1_min or x == x1_max)):
                r, g, b = 0, 190, 255
            elif x1_min < x < x1_max and y1_min < y < y1_max:
                r, g, b = 25, 36, 48

            x2_min, x2_max = int(width * 0.52), int(width * 0.92)
            y2_min, y2_max = int(height * 0.12), int(height * 0.52)
            if (x2_min <= x <= x2_max and (y == y2_min or y == y2_max)) or \
               (y2_min <= y <= y2_max and (x == x2_min or x == x2_max)):
                r, g, b = 0, 190, 255
            elif x2_min < x < x2_max and y2_min < y < y2_max:
                r, g, b = 25, 36, 48

            x3_min, x3_max = int(width * 0.15), int(width * 0.85)
            y3_min, y3_max = int(height * 0.58), int(height * 0.88)
            if (x3_min <= x <= x3_max and (y == y3_min or y == y3_max)) or \
               (y3_min <= y <= y3_max and (x == x3_min or x == x3_max)):
                r, g, b = 140, 150, 165
            elif x3_min < x < x3_max and y3_min < y < y3_max:
                r, g, b = 28, 30, 36

            row.extend([r, g, b])
        rows.append(bytes(row))

    raw_data = b"".join(rows)
    compressed = zlib.compress(raw_data)

    def make_chunk(chunk_type, data):
        length = struct.pack(">I", len(data))
        crc = struct.pack(">I", zlib.crc32(chunk_type + data) & 0xffffffff)
        return length + chunk_type + data + crc

    png_header = b"\x89PNG\r\n\x1a\n"
    ihdr = struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0)
    ihdr_chunk = make_chunk(b"IHDR", ihdr)
    idat_chunk = make_chunk(b"IDAT", compressed)
    iend_chunk = make_chunk(b"IEND", b"")

    os.makedirs(os.path.dirname(filepath), exist_ok=True)
    with open(filepath, "wb") as f:
        f.write(png_header + ihdr_chunk + idat_chunk + iend_chunk)


# ------------------------------------------------------------------
# 8. FastAPI 路由初始化与静态资源挂载
# ------------------------------------------------------------------
app = FastAPI(title="实验室设备数字孪生与不可篡改履历系统")


@app.on_event("startup")
async def init_data():
    check_and_migrate_db()
    os.makedirs("static/assets", exist_ok=True)
    
    # 启动后台硬件心跳断电巡检检测任务
    asyncio.create_task(heartbeat_checker_loop())
    
    target_floorplan = "static/assets/floorplan.png"
    if not os.path.exists(target_floorplan):
        alt_png = "static/assets/大楼实验室分布_01.png"
        if os.path.exists(alt_png):
            shutil.copy(alt_png, target_floorplan)
        elif os.path.exists("static/uploads/default_map.png"):
            shutil.copy("static/uploads/default_map.png", target_floorplan)
        else:
            generate_default_cad_floorplan(target_floorplan)

    db = SessionLocal()

    if db.query(RoomModel).count() == 0:
        rooms = [
            RoomModel(name="302", pos_x="70%", pos_y="23%"),
            RoomModel(name="701", pos_x="63%", pos_y="11%"),
            RoomModel(name="203", pos_x="25%", pos_y="47%"),
        ]
        db.add_all(rooms)
        db.commit()

    if db.query(DeviceModel).count() == 0:
        devs = [
            DeviceModel(
                device_code="DEV-001",
                name="测试设备 1",
                room_name="302",
                status=DeviceStatus.FAULT.value,
                pos_x="20%",
                pos_y="25%",
            ),
            DeviceModel(
                device_code="DEV-002",
                name="测试设备 2",
                room_name="701",
                status=DeviceStatus.FAULT.value,
                pos_x="65%",
                pos_y="25%",
            ),
            DeviceModel(
                device_code="DEV-003",
                name="测试设备 3",
                room_name="302",
                status=DeviceStatus.FAULT.value,
                pos_x="25%",
                pos_y="65%",
            ),
            DeviceModel(
                device_code="DEV-004",
                name="测试设备 4",
                room_name="302",
                status=DeviceStatus.FAULT.value,
                pos_x="70%",
                pos_y="65%",
            ),
        ]
        db.add_all(devs)
        db.commit()

        now = datetime.now()
        logs = [
            DeviceUsageLogModel(
                device_code="DEV-001",
                previous_status="IDLE",
                new_status="IN_USE",
                operator="张博士",
                source_type="MANUAL_WEB",
                raw_data=json.dumps({"action": "start", "current": 1.5}),
                timestamp=now - timedelta(days=2, hours=5),
            ),
        ]
        db.add_all(logs)
        db.commit()
    db.close()


@app.get("/api/v1/hardware/status")
def get_hardware_status():
    """获取全量 ESP32 采样硬件在线 / 断电离线状态列表"""
    db = SessionLocal()
    try:
        devices = db.query(DeviceModel).all()
        now = datetime.now()
        res = []
        for d in devices:
            last_seen = device_last_seen.get(d.device_code)
            is_online = bool(last_seen and (now - last_seen).total_seconds() <= 10)
            res.append({
                "device_code": d.device_code,
                "name": d.name,
                "room_name": d.room_name or "302",
                "status": d.status,
                "is_online": is_online,
                "last_seen": last_seen.strftime("%Y-%m-%d %H:%M:%S") if last_seen else "无心跳记录",
                "seconds_since_heartbeat": round((now - last_seen).total_seconds(), 1) if last_seen else None
            })
        return res
    finally:
        db.close()


@app.get("/api/v1/hardware/logs")
def get_hardware_logs():
    """获取系统全量历史审计与断电感知日志列表（供审计履历报表统一使用）"""
    db = SessionLocal()
    try:
        logs = (
            db.query(DeviceUsageLogModel)
            .order_by(DeviceUsageLogModel.timestamp.desc())
            .limit(300)
            .all()
        )
        all_devs = db.query(DeviceModel).all()
        dev_map = {d.device_code: d for d in all_devs}
        res = []
        for l in logs:
            dev = dev_map.get(l.device_code)
            res.append({
                "id": l.id,
                "timestamp": l.timestamp.strftime("%Y-%m-%d %H:%M:%S") if l.timestamp else "-",
                "device_code": l.device_code,
                "name": dev.name if dev else f"硬件_{l.device_code}",
                "room_name": dev.room_name if dev else "未分配实验室",
                "previous_status": l.previous_status,
                "new_status": l.new_status,
                "operator": l.operator,
                "source_type": l.source_type,
                "raw_data": l.raw_data or "-"
            })
        return res
    finally:
        db.close()


# ------------------------------------------------------------------
# 9. HTTP REST 与 WebSocket 接口
# ------------------------------------------------------------------
@app.websocket("/ws/devices")
async def websocket_devices_endpoint(websocket: WebSocket):
    await manager.connect(websocket)
    try:
        db = SessionLocal()
        try:
            devices = db.query(DeviceModel).all()
            snapshot = [
                {
                    "id": d.id,
                    "device_code": d.device_code,
                    "name": d.name,
                    "room_name": d.room_name or "302",
                    "status": d.status,
                    "pos_x": d.pos_x,
                    "pos_y": d.pos_y,
                }
                for d in devices
            ]
            await websocket.send_json({
                "type": "init",
                "timestamp": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                "devices": snapshot,
            })
        finally:
            db.close()

        while True:
            msg = await websocket.receive_text()
            if msg == "ping":
                await websocket.send_json({"type": "pong"})
    except WebSocketDisconnect:
        manager.disconnect(websocket)
    except Exception:
        manager.disconnect(websocket)


@app.get("/api/v1/building/config")
def get_building_config():
    """获取本地静态 CAD 图路径及每个实验室房间的可拖动百分比坐标 (pos_x, pos_y)"""
    db = SessionLocal()
    try:
        rooms = db.query(RoomModel).all()
        rooms_cfg = [
            {
                "name": r.name,
                "pos_x": r.pos_x or "50%",
                "pos_y": r.pos_y or "50%",
                "color": "#60cdff"
            }
            for r in rooms
        ]
        return {
            "map_url": "/static/assets/floorplan.png",
            "rooms_config": rooms_cfg
        }
    finally:
        db.close()


@app.get("/api/v1/building/map")
def get_building_map():
    return {"map_url": "/static/assets/floorplan.png"}


@app.put("/api/v1/rooms/{room_name}/position")
async def update_room_position(room_name: str, req: UpdateRoomPositionRequest):
    """【拖拽保存：更新实验室房间在 CAD 地图上的 2D 坐标】"""
    db = SessionLocal()
    try:
        room = db.query(RoomModel).filter(RoomModel.name == room_name).first()
        if not room:
            room = RoomModel(name=room_name, pos_x=req.pos_x.strip(), pos_y=req.pos_y.strip())
            db.add(room)
        else:
            room.pos_x = req.pos_x.strip()
            room.pos_y = req.pos_y.strip()

        # 同步关联设备
        devices_in_room = db.query(DeviceModel).filter(DeviceModel.room_name == room_name).all()
        for dev in devices_in_room:
            dev.pos_x = req.pos_x.strip()
            dev.pos_y = req.pos_y.strip()

        db.commit()
        db.refresh(room)

        # 广播房间位置改变事件
        await manager.broadcast({
            "type": "room_position_updated",
            "room_name": room_name,
            "pos_x": room.pos_x,
            "pos_y": room.pos_y,
            "timestamp": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        })

        return {"message": "实验室胶囊位置已保存", "room_name": room_name, "pos_x": room.pos_x, "pos_y": room.pos_y}
    finally:
        db.close()


@app.put("/api/v1/rooms/{old_room_name}")
async def update_room_name(old_room_name: str, req: UpdateRoomNameRequest):
    """【资产管理：重命名实验室房间 API】"""
    db = SessionLocal()
    try:
        new_name_clean = req.new_name.strip()
        if not new_name_clean:
            raise HTTPException(status_code=400, detail="新房间名称不能为空")

        room = db.query(RoomModel).filter(RoomModel.name == old_room_name).first()
        if not room:
            raise HTTPException(status_code=404, detail="源实验室房间不存在")

        if new_name_clean != old_room_name:
            existing = db.query(RoomModel).filter(RoomModel.name == new_name_clean).first()
            if existing:
                raise HTTPException(status_code=400, detail="该房间名称已存在")

            room.name = new_name_clean
            
            devices_in_room = db.query(DeviceModel).filter(DeviceModel.room_name == old_room_name).all()
            for dev in devices_in_room:
                dev.room_name = new_name_clean

            db.commit()

            await manager.broadcast({
                "type": "room_renamed",
                "old_name": old_room_name,
                "new_name": new_name_clean,
                "timestamp": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            })

        return {"message": "重命名实验室成功", "old_name": old_room_name, "new_name": new_name_clean}
    finally:
        db.close()


@app.post("/api/v1/telemetry/ingest")
async def ingest_telemetry(payload: TelemetryIngestPayload):
    db = SessionLocal()
    try:
        results = []
        if payload.channels is not None:
            for ch in payload.channels:
                op = ch.operator or payload.operator or "SYSTEM"
                src = ch.source_type or payload.source_type or "PLC_MODBUS"
                
                raw_input = ch.raw_data
                if isinstance(raw_input, str):
                    raw = {"info": raw_input}
                elif isinstance(raw_input, dict):
                    raw = dict(raw_input)
                else:
                    raw = {}

                if ch.current is not None:
                    raw["current"] = ch.current
                if ch.voltage is not None:
                    raw["voltage"] = ch.voltage

                single_ch = TelemetryChannel(
                    device_code=ch.device_code,
                    status=ch.status,
                    state=ch.state,
                    current=ch.current,
                    voltage=ch.voltage,
                    operator=op,
                    source_type=src,
                    raw_data=raw
                )
                res = await process_telemetry_ingest_single(db, single_ch)
                results.append(res)
        else:
            if not payload.device_code:
                raise HTTPException(status_code=400, detail="单通道模式下 device_code 为必填项")
            
            raw_input = payload.raw_data
            if isinstance(raw_input, str):
                raw = {"info": raw_input}
            elif isinstance(raw_input, dict):
                raw = dict(raw_input)
            else:
                raw = {}

            if payload.current is not None:
                raw["current"] = payload.current
            if payload.voltage is not None:
                raw["voltage"] = payload.voltage

            single_ch = TelemetryChannel(
                device_code=payload.device_code,
                status=payload.status,
                state=payload.state,
                current=payload.current,
                voltage=payload.voltage,
                operator=payload.operator or "SYSTEM",
                source_type=payload.source_type or "MANUAL_WEB",
                raw_data=raw
            )
            res = await process_telemetry_ingest_single(db, single_ch)
            results.append(res)

        if payload.channels is not None:
            return {"status": "success", "count": len(results), "processed": results}
        else:
            res_single = results[0]
            return {
                "status": "success",
                "device_code": res_single["device_code"],
                "previous_status": res_single["previous_status"],
                "new_status": res_single["new_status"],
                "room_name": res_single["room_name"],
                "source_type": res_single["source_type"],
            }
    except HTTPException as he:
        raise he
    except Exception as e:
        db.rollback()
        raise HTTPException(status_code=500, detail=str(e))
    finally:
        db.close()


@app.post("/api/devices/{device_code}/status")
async def change_device_status(device_code: str, req: StatusChangeRequest):
    db = SessionLocal()
    try:
        single_ch = TelemetryChannel(
            device_code=device_code,
            status=req.new_status,
            operator=req.operator or "MANUAL_WEB",
            source_type="MANUAL_WEB",
            raw_data={"action": "manual_click"}
        )
        return await process_telemetry_ingest_single(db, single_ch)
    finally:
        db.close()


@app.get("/api/devices")
def get_all_devices():
    db = SessionLocal()
    try:
        return db.query(DeviceModel).all()
    finally:
        db.close()


@app.post("/api/devices")
async def create_device(req: CreateDeviceRequest):
    db = SessionLocal()
    try:
        existing = db.query(DeviceModel).filter(DeviceModel.device_code == req.device_code).first()
        if existing:
            raise HTTPException(status_code=400, detail="该设备编号已注册")

        count = db.query(DeviceModel).count()
        x_coord = req.pos_x or f"{(15 + (count * 25) % 70)}%"
        y_coord = req.pos_y or f"{(20 + (count * 18) % 60)}%"

        new_dev = DeviceModel(
            device_code=req.device_code.strip(),
            name=req.name.strip(),
            room_name=req.room_name.strip() if req.room_name else "302",
            status=DeviceStatus.IDLE.value,
            pos_x=x_coord,
            pos_y=y_coord,
        )
        db.add(new_dev)
        db.commit()
        db.refresh(new_dev)

        await manager.broadcast({
            "type": "device_created",
            "device": {
                "id": new_dev.id,
                "device_code": new_dev.device_code,
                "name": new_dev.name,
                "room_name": new_dev.room_name,
                "status": new_dev.status,
                "pos_x": new_dev.pos_x,
                "pos_y": new_dev.pos_y,
            },
            "timestamp": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        })

        return {"message": "创建成功", "device": new_dev}
    finally:
        db.close()


@app.get("/api/devices/{device_code}/logs")
def get_device_logs(device_code: str):
    db = SessionLocal()
    try:
        return (
            db.query(DeviceUsageLogModel)
            .filter(DeviceUsageLogModel.device_code == device_code)
            .order_by(DeviceUsageLogModel.timestamp.desc())
            .all()
        )
    finally:
        db.close()


@app.get("/api/v1/rooms")
def get_all_rooms():
    db = SessionLocal()
    try:
        rooms = db.query(RoomModel).all()
        return [r.name for r in rooms]
    finally:
        db.close()


@app.get("/api/v1/rooms/details")
def get_all_rooms_details():
    db = SessionLocal()
    try:
        rooms = db.query(RoomModel).all()
        return [
            {
                "id": r.id,
                "name": r.name,
                "pos_x": r.pos_x or "50%",
                "pos_y": r.pos_y or "50%",
            }
            for r in rooms
        ]
    finally:
        db.close()


@app.post("/api/v1/rooms")
def create_room(req: CreateRoomRequest):
    db = SessionLocal()
    try:
        name_clean = req.name.strip()
        if not name_clean:
            raise HTTPException(status_code=400, detail="房间名称不能为空")
        existing = db.query(RoomModel).filter(RoomModel.name == name_clean).first()
        if existing:
            raise HTTPException(status_code=400, detail="该房间已存在")
        new_room = RoomModel(name=name_clean, pos_x="50%", pos_y="50%")
        db.add(new_room)
        db.commit()
        return {"message": "创建房间成功", "room_name": new_room.name}
    finally:
        db.close()


@app.delete("/api/v1/rooms/{room_name}")
async def delete_room(room_name: str):
    db = SessionLocal()
    try:
        room = db.query(RoomModel).filter(RoomModel.name == room_name).first()
        if not room:
            raise HTTPException(status_code=404, detail="房间不存在")

        fallback_name = "未分配实验室" if room_name != "未分配实验室" else "待划分实验室"
        if room_name == "未分配实验室":
            fallback = db.query(RoomModel).filter(RoomModel.name == fallback_name).first()
            if not fallback:
                fallback = RoomModel(name=fallback_name, pos_x="50%", pos_y="50%")
                db.add(fallback)
                db.commit()

        devices_in_room = db.query(DeviceModel).filter(DeviceModel.room_name == room_name).all()
        for dev in devices_in_room:
            dev.room_name = fallback_name

        db.delete(room)
        db.commit()

        await manager.broadcast({
            "type": "room_deleted",
            "room_name": room_name,
            "timestamp": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        })
        return {"message": f"删除房间 '{room_name}' 成功"}
    finally:
        db.close()


@app.put("/api/devices/{device_code}")
async def update_device(device_code: str, req: UpdateDeviceRequest):
    db = SessionLocal()
    try:
        dev = db.query(DeviceModel).filter(DeviceModel.device_code == device_code).first()
        if not dev:
            raise HTTPException(status_code=404, detail="未找到对应的设备")

        dev.name = req.name.strip()
        dev.room_name = req.room_name.strip()
        if req.pos_x is not None:
            dev.pos_x = req.pos_x.strip()
        if req.pos_y is not None:
            dev.pos_y = req.pos_y.strip()

        db.commit()
        db.refresh(dev)

        await manager.broadcast({
            "type": "device_updated",
            "device": {
                "id": dev.id,
                "device_code": dev.device_code,
                "name": dev.name,
                "room_name": dev.room_name,
                "status": dev.status,
                "pos_x": dev.pos_x,
                "pos_y": dev.pos_y,
            },
            "timestamp": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        })
        return {"message": "编辑设备成功", "device": dev}
    finally:
        db.close()


@app.delete("/api/devices/{device_code}")
async def delete_device(device_code: str):
    db = SessionLocal()
    try:
        dev = db.query(DeviceModel).filter(DeviceModel.device_code == device_code).first()
        if not dev:
            raise HTTPException(status_code=404, detail="未找到对应的设备")

        db.delete(dev)
        db.commit()

        await manager.broadcast({
            "type": "device_deleted",
            "device_code": device_code,
            "timestamp": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        })
        return {"message": "删除设备成功"}
    finally:
        db.close()


@app.get("/api/reports/usage")
def get_usage_reports():
    db = SessionLocal()
    try:
        devices = db.query(DeviceModel).all()
        all_sessions = []
        for dev in devices:
            sessions = calculate_device_usage_sessions(db, dev.device_code, dev.name, dev.room_name)
            for s in sessions:
                s_copy = s.copy()
                s_copy["start_time"] = s_copy["start_time"].strftime("%Y-%m-%d %H:%M:%S")
                if isinstance(s_copy["end_time"], datetime):
                    s_copy["end_time"] = s_copy["end_time"].strftime("%Y-%m-%d %H:%M:%S")
                all_sessions.append(s_copy)
        all_sessions.sort(key=lambda x: x["start_time"], reverse=True)
        return all_sessions
    finally:
        db.close()


@app.get("/api/reports/export-csv")
def export_usage_csv():
    db = SessionLocal()
    try:
        devices = db.query(DeviceModel).all()
        output = io.StringIO()
        output.write("\ufeff")
        writer = csv.writer(output)
        writer.writerow(["设备编号", "设备名称", "所属实验室", "当前状态", "累计使用次数", "累计运行时长(秒)", "格式化时长"])
        
        status_map = {"IDLE": "待机", "IN_USE": "运行中", "FAULT": "故障"}
        for dev in devices:
            sessions = calculate_device_usage_sessions(db, dev.device_code, dev.name, dev.room_name)
            total_sec = sum(s["duration"] for s in sessions)
            writer.writerow([
                dev.device_code,
                dev.name,
                dev.room_name,
                status_map.get(dev.status, dev.status),
                len(sessions),
                round(total_sec, 1),
                format_duration(total_sec)
            ])
            
        csv_content = output.getvalue()
        output.close()
        filename = f"lab_device_csv_{datetime.now().strftime('%Y%m%d_%H%M%S')}.csv"
        return Response(
            content=csv_content,
            media_type="text/csv",
            headers={"Content-Disposition": f"attachment; filename={filename}"},
        )
    finally:
        db.close()


@app.get("/api/reports/export-excel")
def export_usage_excel():
    db = SessionLocal()
    try:
        devices = db.query(DeviceModel).all()
        
        overview_data = []
        for dev in devices:
            sessions = calculate_device_usage_sessions(db, dev.device_code, dev.name, dev.room_name)
            total_sec = sum(s["duration"] for s in sessions)
            
            completed = [s for s in sessions if s["status"] == "COMPLETED"]
            last_used_str = "无记录"
            if completed:
                completed.sort(key=lambda x: x["start_time"])
                last_used_str = completed[-1]["start_time"].strftime("%Y-%m-%d %H:%M:%S")
            elif sessions:
                last_used_str = sessions[0]["start_time"].strftime("%Y-%m-%d %H:%M:%S") + " (运行中)"
                
            overview_data.append({
                "设备编号": dev.device_code,
                "设备名称": dev.name,
                "所属实验室": dev.room_name,
                "当前状态": "运行中" if dev.status == "IN_USE" else ("待机" if dev.status == "IDLE" else "故障"),
                "累计使用次数": len(sessions),
                "累计运行时长": format_duration(total_sec),
                "累计秒数": round(total_sec, 1),
                "最近使用时间": last_used_str
            })
        df_overview = pd.DataFrame(overview_data)

        sessions_data = []
        for dev in devices:
            sessions = calculate_device_usage_sessions(db, dev.device_code, dev.name, dev.room_name)
            for s in sessions:
                end_t = s["end_time"].strftime("%Y-%m-%d %H:%M:%S") if isinstance(s["end_time"], datetime) else "运行中"
                sessions_data.append({
                    "设备编号": s["device_code"],
                    "设备名称": s["name"],
                    "所属实验室": s["room_name"],
                    "启动时间": s["start_time"].strftime("%Y-%m-%d %H:%M:%S"),
                    "结束时间": end_t,
                    "单次运行时长": s["formatted_duration"],
                    "运行时长(秒)": round(s["duration"], 1),
                    "操作人": s["operator"],
                    "启动通道": s["source_type"],
                    "启动遥测参数": s["start_raw"] or "-",
                    "结束遥测参数": s["end_raw"] or "-"
                })
        df_sessions = pd.DataFrame(sessions_data)
        if not df_sessions.empty:
            df_sessions = df_sessions.sort_values(by="启动时间", ascending=False)
        else:
            df_sessions = pd.DataFrame(columns=[
                "设备编号", "设备名称", "所属实验室", "启动时间", "结束时间", "单次运行时长", "运行时长(秒)", "操作人", "启动通道", "启动遥测参数", "结束遥测参数"
            ])

        logs = db.query(DeviceUsageLogModel).order_by(DeviceUsageLogModel.timestamp.desc()).all()
        logs_data = []
        for log in logs:
            logs_data.append({
                "时间": log.timestamp.strftime("%Y-%m-%d %H:%M:%S"),
                "设备编号": log.device_code,
                "变更前状态": "运行中" if log.previous_status == "IN_USE" else ("待机" if log.previous_status == "IDLE" else "故障"),
                "变更后状态": "运行中" if log.new_status == "IN_USE" else ("待机" if log.new_status == "IDLE" else "故障"),
                "操作人": log.operator,
                "来源通道": log.source_type,
                "原始载荷": log.raw_data or "-"
            })
        df_logs = pd.DataFrame(logs_data)
        if df_logs.empty:
            df_logs = pd.DataFrame(columns=["时间", "设备编号", "变更前状态", "变更后状态", "操作人", "来源通道", "原始载荷"])

        output = io.BytesIO()
        with pd.ExcelWriter(output, engine="openpyxl") as writer:
            df_overview.to_excel(writer, sheet_name="设备汇总", index=False)
            df_sessions.to_excel(writer, sheet_name="使用明细", index=False)
            df_logs.to_excel(writer, sheet_name="审计日志", index=False)
            
            workbook = writer.book
            
            from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
            header_fill = PatternFill(start_color="1c1c1c", end_color="1c1c1c", fill_type="solid")
            header_font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
            cell_font = Font(name="Segoe UI", size=10, color="000000")
            grid_border = Border(
                left=Side(style='thin', color='EAEAEA'),
                right=Side(style='thin', color='EAEAEA'),
                top=Side(style='thin', color='EAEAEA'),
                bottom=Side(style='thin', color='EAEAEA')
            )
            align_center = Alignment(horizontal="center", vertical="center")
            align_left = Alignment(horizontal="left", vertical="center")

            for name in workbook.sheetnames:
                sheet = workbook[name]
                sheet.views.sheetView[0].showGridLines = True
                
                for cell in sheet[1]:
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.alignment = align_center
                    cell.border = grid_border
                    
                for row in sheet.iter_rows(min_row=2):
                    for cell in row:
                        cell.font = cell_font
                        cell.border = grid_border
                        
                        if cell.column_letter in ['A', 'D', 'E', 'G', 'H'] or name == "审计日志":
                            cell.alignment = align_center
                        else:
                            cell.alignment = align_left

                for col in sheet.columns:
                    max_len = 0
                    for cell in col:
                        val = str(cell.value or '')
                        width = sum(2 if ord(c) > 255 else 1 for c in val)
                        if width > max_len:
                            max_len = width
                    col_letter = col[0].column_letter
                    sheet.column_dimensions[col_letter].width = max(max_len + 4, 12)

        excel_data = output.getvalue()
        output.close()
        filename = f"lab_twin_history_report_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        return Response(
            content=excel_data,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f"attachment; filename={filename}"},
        )
    finally:
        db.close()


import sys
import threading
import webbrowser

# 动态定位静态资源根目录（兼容源码运行与 PyInstaller _MEIPASS 打包环境）
BASE_DIR = getattr(sys, '_MEIPASS', os.path.dirname(os.path.abspath(__file__)))
STATIC_DIR = os.path.join(BASE_DIR, "static")
if not os.path.exists(STATIC_DIR):
    STATIC_DIR = os.path.abspath("static")

app.mount("/static", StaticFiles(directory=STATIC_DIR), name="static_dir")
app.mount("/", StaticFiles(directory=STATIC_DIR, html=True), name="static")

def open_browser_delayed():
    import time
    time.sleep(1.2)
    try:
        webbrowser.open("http://localhost:8000")
    except Exception:
        pass

if __name__ == "__main__":
    import uvicorn
    # 启动后台线程自动打开浏览器，避免用户误认为服务卡死
    threading.Thread(target=open_browser_delayed, daemon=True).start()
    
    if getattr(sys, 'frozen', False):
        uvicorn.run(app, host="0.0.0.0", port=8000, ws="websockets")
    else:
        uvicorn.run("main:app", host="0.0.0.0", port=8000, reload=True, ws="websockets")